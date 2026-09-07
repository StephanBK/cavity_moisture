"""The surface-energy-balance path, end to end.

The engine gained solar and sky physics in ANLY-002 R1.3 but nothing supplied
it with data. This covers the wiring: NSRDB attributes, WeatherYear fields, the
API parameters, the four-facade comparison, and the Excel sheet.

THE CONTRACT
    Defaults are OFF. A request that does not ask for sun must return exactly
    what it returned before this existed - 860 condensing hours and 676 under
    the room-dew-point criterion at 277 Park. Same precedent as u_ip: never
    silently reinterpret an existing request.

Sources: [LIVE]     NSRDB GOES TMY v4.0.0 with irradiance, wind and cloud
          [API]      developer.nlr.gov, September 2026
          [ANLY-002] anchor preservation
          [PHYSICS]  sun dries, sky cools, cloud suppresses cooling

Doc ID: ANLY-002 R1.3
"""

from __future__ import annotations

import gzip
import io
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.report import workbook_bytes
from engine.solar import ORIENTATIONS
from engine.weather import NSRDB_ATTRIBUTES, Location, parse_nsrdb_csv

FIXTURE = Path(__file__).parent / "fixtures" / "nsrdb_277park_tmy_v4.csv.gz"
BASE_Q = ("address=x&f_cold=0.300&u_ip=0.30&t_in=70&rh_in=35&ach=5"
          "&vent_interior=1&width_in=60&height_in=96&offset_in=0.6024")


@pytest.fixture(scope="module")
def year():
    with gzip.open(FIXTURE, "rt") as fh:
        return parse_nsrdb_csv(fh.read())


@pytest.fixture
def client(monkeypatch, year):
    os.environ.setdefault("MAPBOX_TOKEN", "pk.test")
    os.environ.setdefault("NREL_API_KEY", "test")
    import app as flask_app

    loc = Location("x", "277 PARK AVE, NEW YORK, NY 10172", 40.7568, -73.9742)
    monkeypatch.setattr(
        flask_app, "get_weather_for_address", lambda *a, **k: (loc, year, False)
    )
    return flask_app.app.test_client()


# ---------------------------------------------------------------------------
# THE ANCHOR  [ANLY-002]
# ---------------------------------------------------------------------------

def test_defaults_reproduce_the_anchor(client):
    """No solar parameters means the air-only model, unchanged."""
    s = client.get("/calculate?" + BASE_Q).get_json()["summary"]
    assert s["hours_condensing"] == 860
    assert s["hours_condensing_room_assumption"] == 676
    assert s["condensed_g_per_m2"] == pytest.approx(53.042, abs=0.01)


def test_absorptance_zero_is_the_same_as_omitting_it(client):
    a = client.get("/calculate?" + BASE_Q).get_json()["summary"]
    b = client.get("/calculate?" + BASE_Q + "&absorptance=0").get_json()["summary"]
    assert a["hours_condensing"] == b["hours_condensing"]


# ---------------------------------------------------------------------------
# The fetch  [API] [LIVE]
# ---------------------------------------------------------------------------

def test_the_fetch_requests_everything_the_balance_needs():
    for attr in ("air_temperature", "relative_humidity", "surface_pressure",
                 "ghi", "dni", "dhi", "wind_speed", "cloud_type"):
        assert attr in NSRDB_ATTRIBUTES, attr


def test_the_fixture_carries_a_full_solar_year(year):
    assert year.hours == 8760
    assert year.has_solar
    for series in (year.ghi_w_m2, year.dni_w_m2, year.dhi_w_m2,
                   year.wind_m_s, year.cloud_type):
        assert len(series) == 8760


def test_irradiance_is_physically_plausible(year):
    """NYC horizontal totals roughly 1400-1600 kWh/m2/yr. [LIVE]"""
    assert 900 < max(year.ghi_w_m2) < 1200
    assert 1300 < sum(year.ghi_w_m2) / 1000 < 1700
    assert min(year.ghi_w_m2) >= 0.0


def test_the_source_label_names_the_dataset_actually_served(year):
    """It said PSM3 long after that endpoint was retired. [API]"""
    assert "GOES" in year.source and "v4" in year.source
    assert "PSM3" not in year.source


def test_a_weather_year_without_irradiance_is_honest_about_it():
    csv_text = (
        "Source,Location ID,Latitude,Longitude,Time Zone,Elevation\n"
        "NSRDB,1,40.77,-73.98,-5,20\n"
        "Year,Month,Day,Hour,Minute,Temperature,Relative Humidity,Pressure\n"
        + "".join(f"2013,1,1,0,30,5.0,60,1013\n" for _ in range(8760))
    )
    y = parse_nsrdb_csv(csv_text)
    assert not y.has_solar
    assert y.ghi_w_m2 == []


def test_describe_publishes_whether_solar_is_available(year):
    assert year.describe()["has_solar"] is True


# ---------------------------------------------------------------------------
# The parameters  [API]
# ---------------------------------------------------------------------------

def test_sun_reduces_condensation(client):
    """[PHYSICS] Warming the cold face cannot create water."""
    base = client.get("/calculate?" + BASE_Q).get_json()["summary"]
    sunny = client.get(
        "/calculate?" + BASE_Q + "&absorptance=0.10&orientation=south"
    ).get_json()["summary"]
    assert sunny["hours_water_present"] < base["hours_water_present"]


def test_sky_radiation_increases_condensation(client):
    """[PHYSICS] Radiative cooling makes the glass colder than the air."""
    base = client.get("/calculate?" + BASE_Q).get_json()["summary"]
    cold = client.get(
        "/calculate?" + BASE_Q + "&sky_radiation=true"
    ).get_json()["summary"]
    assert cold["hours_water_present"] > base["hours_water_present"]


def test_orientation_changes_the_answer(client):
    q = BASE_Q + "&absorptance=0.20"
    north = client.get(f"/calculate?{q}&orientation=north").get_json()["summary"]
    south = client.get(f"/calculate?{q}&orientation=south").get_json()["summary"]
    assert south["hours_water_present"] < north["hours_water_present"]


def test_an_unknown_orientation_is_rejected(client):
    r = client.get("/calculate?" + BASE_Q + "&orientation=up")
    assert r.status_code == 400
    assert "orientation" in r.get_json()["message"]


def test_absorptance_outside_zero_to_one_is_rejected(client):
    assert client.get("/calculate?" + BASE_Q + "&absorptance=1.5").status_code == 400


def test_inputs_echo_the_surface_balance_settings(client):
    got = client.get(
        "/calculate?" + BASE_Q + "&absorptance=0.10&sky_radiation=true&orientation=east"
    ).get_json()["inputs"]
    assert got["absorptance"] == pytest.approx(0.10)
    assert got["sky_radiation"] is True
    assert got["orientation"] == "east"


# ---------------------------------------------------------------------------
# The four-facade comparison  [PHYSICS]
# ---------------------------------------------------------------------------

@pytest.fixture
def facades(client):
    d = client.get(
        "/calculate?" + BASE_Q
        + "&absorptance=0.10&sky_radiation=true&orientations=true"
    ).get_json()
    return {r["orientation"]: r for r in d["orientations"]}


def test_all_four_facades_are_reported(facades):
    assert set(facades) == set(ORIENTATIONS)


def test_south_is_the_driest_and_north_the_wettest(facades):
    """Low winter sun strikes vertical south glass near normal, in exactly
    the season condensation peaks."""
    order = sorted(facades, key=lambda k: facades[k]["hours_water_visible"])
    assert order[0] == "south"
    assert order[-1] == "north"


def test_the_facade_spread_is_material(facades):
    """If this collapses, the feature is not earning its complexity."""
    hours = [f["hours_water_visible"] for f in facades.values()]
    assert (max(hours) - min(hours)) / max(hours) > 0.15


def test_east_and_west_sit_between_north_and_south(facades):
    for side in ("east", "west"):
        assert (facades["south"]["hours_water_visible"]
                < facades[side]["hours_water_visible"]
                < facades["north"]["hours_water_visible"])


def test_no_facade_panel_without_sun(client):
    """With absorptance 0 every facade is identical, and publishing four equal
    rows would imply orientation does not matter rather than that it was not
    modelled."""
    d = client.get("/calculate?" + BASE_Q + "&orientations=true").get_json()
    assert "orientations" not in d


# ---------------------------------------------------------------------------
# Excel  [SPEC]
# ---------------------------------------------------------------------------

def test_the_workbook_gains_an_orientation_sheet(year):
    from engine.cavity import f_warm_estimate
    from engine.moisture import run_year
    from engine.psychro import f_to_c

    summary = run_year(
        year.t_out_c, year.rh_out, f_cold=0.300,
        f_warm=f_warm_estimate(0.300, 0.17, 1.7), ach=5.0,
        t_room_c=f_to_c(70), rh_room=0.35, vent_interior_fraction=1.0,
        elevation_m=year.elevation_m,
    )
    rows = [{"orientation": n, "azimuth_deg": az, "hours_water_visible": 100,
             "hours_water_present": 200, "condensed_g_per_m2": 50.0,
             "per_window": {"condensed_litres_per_year": 0.2}}
            for n, az in ORIENTATIONS.items()]
    wb = load_workbook(io.BytesIO(
        workbook_bytes(summary, {"address": "t"}, [summary], ["moderate"], rows)
    ))
    assert "Orientation" in wb.sheetnames
    ws = wb["Orientation"]
    assert [ws.cell(row=r, column=1).value for r in range(2, 6)] == list(ORIENTATIONS)


def test_no_orientation_sheet_when_none_supplied(year):
    from engine.cavity import f_warm_estimate
    from engine.moisture import run_year
    from engine.psychro import f_to_c

    summary = run_year(
        year.t_out_c, year.rh_out, f_cold=0.300,
        f_warm=f_warm_estimate(0.300, 0.17, 1.7), ach=5.0,
        t_room_c=f_to_c(70), rh_room=0.35, vent_interior_fraction=1.0,
        elevation_m=year.elevation_m,
    )
    wb = load_workbook(io.BytesIO(
        workbook_bytes(summary, {"address": "t"}, [summary], ["moderate"])
    ))
    assert "Orientation" not in wb.sheetnames


# ---------------------------------------------------------------------------
# The frontend contract  [A11Y]
# ---------------------------------------------------------------------------

def test_the_page_sends_the_new_parameters():
    html = Path(__file__).resolve().parents[1].joinpath("static/index.html").read_text()
    for name in ("orientation", "absorptance", "sky_radiation", "orientations"):
        assert name in html, f"page never sends {name!r}"


def test_the_page_labels_absorptance_as_an_estimate():
    html = Path(__file__).resolve().parents[1].joinpath("static/index.html").read_text()
    assert "ESTIMATE" in html
